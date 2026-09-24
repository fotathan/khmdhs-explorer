"""The checklist on paper and in a spreadsheet (tender-checklist spec, slice 5).

What is defended here:

  SAME LIST   the print page and the .xlsx are built from view(), the panel's
              own dict: same items, same ticks, same own items, same counts.
  PRIVATE     a customer's export carries THEIR ticks and own items only; the
              response is no-store and noindex.
  DATED       both say when they were made — a printed "5 days left" is
              stated from a date, because paper does not update itself.
  WARNED      the screening warning travels with the list.
  GATED       anonymous, non-entitled, AI-off and no-checklist requests are
              sent to the act page, never shown the list or an error.
  NO FORMULAS every text cell in the workbook is a string: notice quotes and
              customer text starting with "=" must never become formulas.
"""
import datetime as dt
import io

import pytest
from openpyxl import load_workbook

from tests.helpers import connect, get_csrf, grant, login, make_user

ADAM = "TEST-CLEXPORT-0001"
TEXT = ("ΔΙΑΚΗΡΥΞΗ\n\n"
        "Η προθεσμία υποβολής ερωτημάτων λήγει στις 12/06/2030 και ώρα 15:00.\n\n"
        "Απαιτείται πιστοποίηση ISO 9001:2015 σε ισχύ.\n\n"
        "Η προσφορά υποβάλλεται ηλεκτρονικά μέσω ΕΣΗΔΗΣ μαζί με το ΕΕΕΣ.\n")


def _item(label, value, quote, obligation="mandatory"):
    return {"label": label, "value": value, "obligation": obligation,
            "confidence": "high", "source": "full_text", "quote": quote}


RAW = {
    "timeline": [_item("Προθεσμία ερωτημάτων", "12/06/2030, 15:00",
                       "Η προθεσμία υποβολής ερωτημάτων λήγει στις 12/06/2030")],
    "eligibility": [_item("Πιστοποίηση ποιότητας", "ISO 9001:2015",
                          "πιστοποίηση ISO 9001:2015 σε ισχύ")],
    "submission": [_item("ΕΕΕΣ", "Μέσω ΕΣΗΔΗΣ",
                         "υποβάλλεται ηλεκτρονικά μέσω ΕΣΗΔΗΣ μαζί με το ΕΕΕΣ",
                         obligation="unspecified")],
    "not_found": [],
}


@pytest.fixture()
def act(db):
    cur = db.cursor()
    cur.execute("DELETE FROM proc.procurement_act WHERE adam=%s", (ADAM,))
    cur.execute("""INSERT INTO proc.procurement_act
                     (adam, type, title, origin, data_source, full_text,
                      final_submission_date)
                   VALUES (%s, 'notice', 'Προμήθεια για εκτύπωση', 'import',
                           'khmdhs', %s, '2030-06-20 10:00+03')""", (ADAM, TEXT))
    yield ADAM
    cur.execute("DELETE FROM proc.act_ai_summary_history WHERE adam=%s", (ADAM,))
    cur.execute("DELETE FROM proc.procurement_act WHERE adam=%s", (ADAM,))


@pytest.fixture()
def ai_on(monkeypatch):
    monkeypatch.setenv("AI_SUMMARY_ENABLED", "1")
    from app import ai_summary as _ai
    monkeypatch.setenv(_ai.key_var(), "test-key-not-used")
    from app.main import templates
    monkeypatch.setitem(templates.env.globals, "ai_summary_enabled", True)


def _store():
    from app import ai_summary as ai
    with connect() as conn:
        cur = conn.cursor()
        act_row, sources = ai.load_inputs(cur, ADAM)
        ai.store(cur, ADAM, payload=ai.verify(RAW, sources, act_row),
                 hash_=ai.input_hash(sources), usage={}, by="test")


def _customer(client, name="clexp", *, entitled=True):
    uid = make_user(name, "goodpassword1")
    if entitled:
        grant(uid)
    login(client, name, "goodpassword1")
    return uid


def _key(section, label, quote):
    from app import tender_checklist as tc
    return tc.item_key(section, label, quote)


def _tick_iso(client):
    key = _key("eligibility", "Πιστοποίηση ποιότητας",
               "πιστοποίηση ISO 9001:2015 σε ισχύ")
    r = client.post(f"/act/{ADAM}/checklist/{key}", data={"done": "1"},
                    headers={"X-CSRF-Token": get_csrf(client)})
    assert r.status_code == 200


def _own(client, text):
    r = client.post(f"/act/{ADAM}/checklist/own", data={"text": text},
                    headers={"X-CSRF-Token": get_csrf(client)})
    assert r.status_code == 200


def _xlsx(client):
    r = client.get(f"/act/{ADAM}/checklist.xlsx")
    assert r.status_code == 200, r.text
    return r, load_workbook(io.BytesIO(r.content))


def _cells(ws):
    return [c for row in ws.iter_rows() for c in row if c.value is not None]


# --------------------------------------------------------------------------- #
# The print page
# --------------------------------------------------------------------------- #
def test_the_print_page_is_the_panels_list(client, act, ai_on):
    _store()
    _customer(client)
    _tick_iso(client)
    _own(client, "Αίτημα εγγυητικής στην τράπεζα")
    r = client.get(f"/act/{ADAM}/checklist/print")
    assert r.status_code == 200
    body = r.text
    assert "Προμήθεια για εκτύπωση" in body and ADAM in body
    assert "Πιστοποίηση ποιότητας" in body and "ΕΕΕΣ" in body
    assert "Αίτημα εγγυητικής στην τράπεζα" in body
    assert "Προθεσμία ερωτημάτων" in body and "12/06/2030 15:00" in body
    assert "20/06/2030 10:00" in body                   # the record, Athens time
    assert "1 / 3" in body                              # 2 AI items + 1 own, 1 ticked
    assert body.count("✓") == 1
    assert "Εργαλείο πρώτης αξιολόγησης." in body        # the warning travels
    assert "οι ημέρες μετρούν από αυτή την ημερομηνία." in body
    assert "noindex" in body
    assert r.headers["x-robots-tag"] == "noindex"
    assert "no-store" in r.headers["cache-control"]


def test_quotes_can_be_left_off(client, act, ai_on):
    _store()
    _customer(client)
    quote = "πιστοποίηση ISO 9001:2015 σε ισχύ"
    assert quote in client.get(f"/act/{ADAM}/checklist/print").text
    body = client.get(f"/act/{ADAM}/checklist/print?quotes=0").text
    assert quote not in body and "Πιστοποίηση ποιότητας" in body


def test_an_export_carries_only_my_ticks_and_items(client, act, ai_on):
    _store()
    _customer(client, "clexp_a")
    _tick_iso(client)
    _own(client, "Μόνο δικό μου")
    client.cookies.clear()
    _customer(client, "clexp_b")
    body = client.get(f"/act/{ADAM}/checklist/print").text
    assert "Μόνο δικό μου" not in body and "✓" not in body
    _r, wb = _xlsx(client)
    values = {c.value for ws in wb for c in _cells(ws)}
    assert "Μόνο δικό μου" not in values


def test_the_panel_links_to_both(client, act, ai_on):
    _store()
    _customer(client)
    body = client.get(f"/act/{ADAM}/checklist").text
    assert f"/act/{ADAM}/checklist/print" in body
    assert f"/act/{ADAM}/checklist.xlsx" in body


# --------------------------------------------------------------------------- #
# The spreadsheet
# --------------------------------------------------------------------------- #
def test_the_workbook_has_the_tasks_and_the_deadlines(client, act, ai_on):
    _store()
    _customer(client)
    _tick_iso(client)
    _own(client, "Βιογραφικά μηχανικών")
    r, wb = _xlsx(client)
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert f'filename="checklist-{ADAM}.xlsx"' in r.headers["content-disposition"]
    assert "no-store" in r.headers["cache-control"]
    assert wb.sheetnames == ["Λίστα ελέγχου", "Προθεσμίες"]

    tasks = wb["Λίστα ελέγχου"]
    rows = [[c.value for c in row] for row in tasks.iter_rows()]
    labels = [r[1] for r in rows]
    assert "Πιστοποίηση ποιότητας" in labels and "Βιογραφικά μηχανικών" in labels
    iso = rows[labels.index("Πιστοποίηση ποιότητας")]
    assert isinstance(iso[4], dt.datetime)             # a real date cell: done
    assert rows[labels.index("ΕΕΕΣ")][4] is None       # not ticked
    flat = " ".join(str(v) for r in rows for v in r if v)
    assert "Η λίστα προκύπτει αυτόματα" in flat

    dates = wb["Προθεσμίες"]
    rows = [[c.value for c in row] for row in dates.iter_rows()]
    by_label = {r[2]: r for r in rows if r[2]}
    assert by_label["Υποβολή προσφορών"][0] == dt.datetime(2030, 6, 20)
    assert by_label["Υποβολή προσφορών"][1] == "10:00"
    assert by_label["Προθεσμία ερωτημάτων"][0] == dt.datetime(2030, 6, 12)
    assert isinstance(by_label["Προθεσμία ερωτημάτων"][4], int)   # days left


def test_the_workbook_follows_the_language(client, act, ai_on):
    _store()
    _customer(client)
    client.cookies.set("lang", "en")
    _r, wb = _xlsx(client)
    assert wb.sheetnames == ["Checklist", "Deadlines"]
    headers = {c.value for c in _cells(wb["Checklist"])}
    assert {"Section", "Item", "What the notice says"} <= headers


def test_no_text_cell_ever_becomes_a_formula(client, act, ai_on):
    _store()
    _customer(client)
    _own(client, '=HYPERLINK("http://evil.example","click")')
    _r, wb = _xlsx(client)
    for ws in wb:
        for c in _cells(ws):
            assert c.data_type != "f", (ws.title, c.coordinate, c.value)
    texts = [c.value for c in _cells(wb["Λίστα ελέγχου"])]
    assert '=HYPERLINK("http://evil.example","click")' in texts


def test_a_greek_ada_keeps_its_name_in_the_download():
    from urllib.parse import quote
    from app import checklist_export as ex
    h = ex.disposition("ΨΔΞΞΟΡΛΟ-Χ5Δ", "xlsx")
    assert h.startswith('attachment; filename="checklist-')       # ASCII fallback
    assert h.split('filename="')[1].split('"')[0].isascii()
    assert "filename*=UTF-8''" + quote("checklist-ΨΔΞΞΟΡΛΟ-Χ5Δ.xlsx", safe="") in h


def test_the_writer_itself_refuses_formulas():
    """Unit level: a quote from the notice starting with '=' stays text."""
    from app import checklist_export as ex
    cl = {"adam": "X", "title": "=1+1", "groups": [{
              "heading": "Υποβολή προσφοράς",
              "items": [{"label": "=SUM(A1:A2)", "value": "=2*3",
                         "quote": "=cmd|' /C calc'!A0", "obligation": None,
                         "done_at": None}]}],
          "own": [], "deadlines": [], "truncated": False}
    meta = {"title": "=1+1", "authority": None,
            "made_at": dt.datetime(2030, 1, 1, 9, 0)}
    wb = load_workbook(io.BytesIO(ex.workbook(cl, meta)))
    for ws in wb:
        for c in _cells(ws):
            assert c.data_type != "f", (ws.title, c.coordinate, c.value)


# --------------------------------------------------------------------------- #
# Gating
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("path", ["checklist/print", "checklist.xlsx"])
def test_anonymous_and_unentitled_go_to_the_act(client, act, ai_on, path):
    _store()
    r = client.get(f"/act/{ADAM}/{path}", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == f"/act/{ADAM}"
    _customer(client, "clexp_free", entitled=False)
    r = client.get(f"/act/{ADAM}/{path}", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == f"/act/{ADAM}"


@pytest.mark.parametrize("path", ["checklist/print", "checklist.xlsx"])
def test_no_checklist_or_ai_off_goes_to_the_act(client, act, ai_on, monkeypatch, path):
    _customer(client)
    r = client.get(f"/act/{ADAM}/{path}", follow_redirects=False)   # no summary
    assert r.status_code == 303 and r.headers["location"] == f"/act/{ADAM}"
    _store()
    monkeypatch.setenv("AI_SUMMARY_ENABLED", "0")
    r = client.get(f"/act/{ADAM}/{path}", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == f"/act/{ADAM}"
