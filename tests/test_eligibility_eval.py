"""The evaluation layer: declared certificates against the checklist
(docs/specs/evaluation-layer.md).

What is defended here:

  MATCHING    the catalogue, in the spellings real Greek notices use (Greek
              lookalike letters, «ΕΝ ISO», lists, the 13458 typo); product
              standards are never matched; OHSAS → 45001 one way only.
  WHOSE       an item naming «ο κατασκευαστής» is answered from the
              manufacturer rows, one naming both from both.
  VALIDITY    ok / lapses before the closing date / expired / no date, on
              Athens dates; an edition mismatch is said, not hidden.
  SUGGESTS    notes never tick anything.
  NEUTRAL     an undeclared scheme is "not declared", never "you lack it" —
              and only for a customer who declared something at all: with no
              certificates the panel is byte-identical to before.
  PRIVATE     user B never sees user A's certificates.
  EVERYWHERE  print, Excel (as a STRING cell) and the favourites line.
  ADMIN ONLY  the CRM routes refuse a customer.
  ISOLATION   ai_summary.py never reads certificates; eligibility_eval.py
              never writes the shared summary.
"""
import datetime as dt
import io
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import eligibility_eval as ev
from tests.helpers import connect, get_csrf, grant, login, make_user

ROOT = Path(__file__).resolve().parent.parent
TODAY = dt.date(2030, 6, 1)
CLOSING = dt.date(2030, 6, 20)


def _schemes(text):
    return [(r["scheme"], r["edition"]) for r in ev.requested(text)]


def _cert(scheme, until=None, *, holder="self", manufacturer=None, edition=None):
    return {"scheme": scheme, "holder": holder, "manufacturer": manufacturer,
            "edition": edition, "valid_until": until}


# --------------------------------------------------------------------------- #
# Matching — written in Greek on purpose: ASCII tests pass broken code
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("text, want", [
    ("Πιστοποιητικό ISO 9001:2015 σε ισχύ", [("iso9001", "2015")]),
    ("πιστοποίηση κατά ΕΛΟΤ ΕΝ ISO 9001", [("iso9001", None)]),
    ("Πιστοποιητικό ISO9001", [("iso9001", None)]),
    ("ΙSΟ 14001 (με ελληνικά Ι και Ο)", [("iso14001", None)]),
    ("ISO 9001:2015 και 14001:2015", [("iso9001", "2015"), ("iso14001", "2015")]),
    ("ISO 9001, 14001 & 45001", [("iso9001", None), ("iso14001", None),
                                 ("iso45001", None)]),
    ("ISO/IEC 27001:2013", [("iso27001", "2013")]),
    ("ISO 13458 για ιατροτεχνολογικά", [("iso13485", None)]),
    ("σύστημα διαχείρισης OHSAS 18001", [("iso45001", None)]),
    ("σύστημα HACCP ή ΕΛΟΤ 1416", [("haccp", None)]),
])
def test_the_catalogue_in_real_spellings(text, want):
    assert _schemes(text) == want


@pytest.mark.parametrize("text", [
    "Τα προϊόντα να φέρουν σήμανση κατά ISO 10993",
    "Επισήμανση σύμφωνα με ISO 15223-1",
    "Συνδέσμοι κατά ISO 7376",
    "Απαιτείται εμπειρία τριών ετών",
    "Έκδοση 2015",
])
def test_product_standards_and_plain_text_are_never_matched(text):
    assert ev.requested(text) == []


def test_a_list_stops_at_a_number_outside_the_catalogue():
    assert _schemes("ISO 9001, 10993, 14001") == [("iso9001", None)]


def test_ohsas_is_answered_by_45001_but_cannot_be_declared():
    notes = ev.evaluate("OHSAS 18001", [_cert("iso45001", CLOSING)],
                        closing=CLOSING, today=TODAY)
    assert [n["status"] for n in notes] == ["ok"]
    assert "ohsas18001" not in ev.SCHEMES              # the reverse is impossible


# --------------------------------------------------------------------------- #
# Whose certificate
# --------------------------------------------------------------------------- #
def test_who_must_hold_it():
    assert ev.holders_for("Ο προσφέρων να διαθέτει ISO 9001") == ("self",)
    assert ev.holders_for("ISO 9001 σε ισχύ") == ("self",)
    assert ev.holders_for("Ο κατασκευαστής να διαθέτει ISO 13485") == ("manufacturer",)
    assert ev.holders_for("Ο οικονομικός φορέας ή ο κατασκευαστής, ISO 13485") == \
        ("self", "manufacturer")


def test_a_manufacturer_item_is_answered_from_manufacturer_rows():
    certs = [_cert("iso13485", CLOSING),
             _cert("iso13485", dt.date(2031, 1, 1), holder="manufacturer",
                   manufacturer="Medtek GmbH")]
    notes = ev.evaluate("Ο κατασκευαστής να διαθέτει ISO 13485", certs,
                        closing=CLOSING, today=TODAY)
    assert [(n["holder"], n["manufacturer"]) for n in notes] == \
        [("manufacturer", "Medtek GmbH")]


def test_a_manufacturer_item_with_none_declared_is_neutral():
    notes = ev.evaluate("Ο κατασκευαστής να διαθέτει ISO 13485",
                        [_cert("iso9001", CLOSING)], closing=CLOSING, today=TODAY)
    assert [(n["status"], n["holder"]) for n in notes] == [("undeclared", "manufacturer")]
    assert "(κατασκευαστή)" in ev.note_text(notes[0])


# --------------------------------------------------------------------------- #
# Validity, on the Athens calendar
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("until, status", [
    (dt.date(2030, 6, 20), "ok"),                 # valid ON the closing day
    (dt.date(2031, 1, 1), "ok"),
    (dt.date(2030, 6, 19), "expires_before"),
    (dt.date(2030, 5, 31), "expired"),
    (None, "no_date"),
])
def test_validity_against_the_closing_date(until, status):
    notes = ev.evaluate("ISO 9001", [_cert("iso9001", until)],
                        closing=CLOSING, today=TODAY)
    assert notes[0]["status"] == status


def test_without_a_closing_date_validity_is_judged_on_today():
    notes = ev.evaluate("ISO 9001", [_cert("iso9001", TODAY)],
                        closing=None, today=TODAY)
    assert notes[0]["status"] == "ok"


def test_an_edition_mismatch_is_said():
    notes = ev.evaluate("ISO 9001:2015", [_cert("iso9001", CLOSING, edition="2008")],
                        closing=CLOSING, today=TODAY)
    assert notes[0]["edition_differs"] is True
    assert "η προκήρυξη αναφέρει έκδοση 2015" in ev.note_text(notes[0])
    same = ev.evaluate("ISO 9001:2015", [_cert("iso9001", CLOSING, edition="2015")],
                       closing=CLOSING, today=TODAY)
    assert same[0]["edition_differs"] is False


def test_several_schemes_in_one_item_get_one_line_each():
    notes = ev.evaluate("ISO 9001 και 14001", [_cert("iso9001", CLOSING)],
                        closing=CLOSING, today=TODAY)
    assert [(n["scheme"], n["status"]) for n in notes] == \
        [("iso9001", "ok"), ("iso14001", "undeclared")]


def test_no_certificates_means_no_notes_at_all():
    groups = [{"items": [{"label": "ISO", "value": "ISO 9001", "quote": "ISO 9001"}]}]
    assert ev.annotate(groups, [], closing=CLOSING, today=TODAY) == 0
    assert "certs" not in groups[0]["items"][0]


def test_the_catalogue_matches_the_migrations_check():
    sql = (ROOT / "migrations" / "20260924150000_company_certificate.sql").read_text()
    listed = set(re.findall(r"'([a-z0-9]+)'", sql.split("scheme IN (", 1)[1].split(")", 1)[0]))
    assert listed == set(ev.SCHEMES)


# --------------------------------------------------------------------------- #
# The checklist, end to end
# --------------------------------------------------------------------------- #
ADAM = "TEST-CERTEVAL-0001"
TEXT = ("ΔΙΑΚΗΡΥΞΗ\n\n"
        "Ο προσφέρων απαιτείται να διαθέτει πιστοποίηση ISO 9001:2015 σε ισχύ.\n\n"
        "Ο οικονομικός φορέας διαθέτει σύστημα ISO 14001.\n\n"
        "Η προσφορά υποβάλλεται ηλεκτρονικά μέσω ΕΣΗΔΗΣ μαζί με το ΕΕΕΣ.\n")
FAR = dt.date.today() + dt.timedelta(days=40)


def _item(label, value, quote, obligation="mandatory"):
    return {"label": label, "value": value, "obligation": obligation,
            "confidence": "high", "source": "full_text", "quote": quote}


RAW = {
    "eligibility": [
        _item("Πιστοποίηση ποιότητας", "ISO 9001:2015",
              "απαιτείται να διαθέτει πιστοποίηση ISO 9001:2015 σε ισχύ"),
        _item("Περιβαλλοντική διαχείριση", "ISO 14001",
              "Ο οικονομικός φορέας διαθέτει σύστημα ISO 14001")],
    "submission": [_item("ΕΕΕΣ", "Μέσω ΕΣΗΔΗΣ",
                         "υποβάλλεται ηλεκτρονικά μέσω ΕΣΗΔΗΣ μαζί με το ΕΕΕΣ",
                         obligation="unspecified")],
    "not_found": [],
}


@pytest.fixture()
def act(db):
    cur = db.cursor()
    cur.execute("DELETE FROM proc.procurement_act WHERE adam=%s", (ADAM,))
    cur.execute("""INSERT INTO proc.procurement_act
                     (adam, type, title, origin, data_source, full_text,
                      final_submission_date)
                   VALUES (%s, 'notice', 'Προμήθεια με πιστοποιητικά', 'import',
                           'khmdhs', %s, %s)""",
                (ADAM, TEXT, dt.datetime.combine(FAR, dt.time(10, 0))))
    yield ADAM
    cur.execute("DELETE FROM proc.act_ai_summary_history WHERE adam=%s", (ADAM,))
    cur.execute("DELETE FROM proc.procurement_act WHERE adam=%s", (ADAM,))


@pytest.fixture()
def ai_on(monkeypatch):
    monkeypatch.setenv("AI_SUMMARY_ENABLED", "1")
    from app import ai_summary as _ai
    monkeypatch.setenv(_ai.key_var(), "test-key-not-used")
    from app.main import templates
    monkeypatch.setitem(templates.env.globals, "ai_summary_enabled", True)


def _store():
    from app import ai_summary as ai
    with connect() as conn:
        cur = conn.cursor()
        act_row, sources = ai.load_inputs(cur, ADAM)
        ai.store(cur, ADAM, payload=ai.verify(RAW, sources, act_row),
                 hash_=ai.input_hash(sources), usage={}, by="test")


def _customer(client, name="certeval", *, favourite=False):
    uid = make_user(name, "goodpassword1")
    grant(uid)
    login(client, name, "goodpassword1")
    if favourite:
        with connect() as conn:
            conn.cursor().execute(
                "INSERT INTO proc.user_favorite_act (user_id, adam, bid_stage) "
                "VALUES (%s, %s, 'bidding')", (uid, ADAM))
    return uid


def _declare(uid, scheme, until, **kw):
    with connect() as conn:
        ev.save(conn.cursor(), uid, scheme=scheme,
                valid_until=until.isoformat() if until else "", **kw)


def test_the_panel_shows_the_profile_under_the_item(client, act, ai_on):
    _store()
    uid = _customer(client)
    _declare(uid, "iso9001", FAR + dt.timedelta(days=200), edition="2015")
    body = client.get(f"/act/{ADAM}/checklist").text
    assert "Στο προφίλ σας" in body and "ISO 9001:2015" in body
    assert "ισχύει έως" in body
    # ISO 14001 was asked for and not declared: neutral, never a verdict.
    assert "ISO 14001" in body and "δεν έχει δηλωθεί στο προφίλ σας." in body
    # Suggests, never ticks.
    assert "checked" not in body
    assert "0 / 3" in body


def test_a_certificate_lapsing_before_the_deadline_is_flagged(client, act, ai_on):
    _store()
    uid = _customer(client, favourite=True)
    _declare(uid, "iso9001", FAR - dt.timedelta(days=3))
    body = client.get(f"/act/{ADAM}/checklist").text
    assert "πριν την υποβολή" in body and "is-expires_before" in body
    fav = client.get("/account/favorites").text
    assert 'class="fav-cl-cert"' in fav and "ISO 9001" in fav
    assert "λήγει πριν την υποβολή" in fav


def test_without_certificates_the_panel_is_unchanged(client, act, ai_on, monkeypatch):
    _store()
    _customer(client)
    body = client.get(f"/act/{ADAM}/checklist").text
    # The same render with the evaluation layer switched off entirely.
    monkeypatch.setattr(ev, "annotate", lambda *a, **k: 0)
    assert client.get(f"/act/{ADAM}/checklist").text == body
    assert "cl-certs" not in body and "Στο προφίλ σας" not in body
    fav_line = client.get("/account/favorites").text
    assert 'class="fav-cl-cert"' not in fav_line


def test_another_customer_never_sees_them(client, act, ai_on):
    _store()
    a = _customer(client, "certa")
    _declare(a, "iso9001", FAR + dt.timedelta(days=200))
    client.cookies.clear()
    _customer(client, "certb")
    body = client.get(f"/act/{ADAM}/checklist").text
    assert "Στο προφίλ σας" not in body and "cl-certs" not in body


def test_print_and_excel_carry_the_notes(client, act, ai_on):
    _store()
    uid = _customer(client)
    _declare(uid, "iso9001", FAR + dt.timedelta(days=200))
    _declare(uid, "iso13485", None, holder="manufacturer", manufacturer="=HYPERLINK(1)")
    printed = client.get(f"/act/{ADAM}/checklist/print").text
    assert "Στο προφίλ σας" in printed and "ISO 9001" in printed
    r = client.get(f"/act/{ADAM}/checklist.xlsx")
    ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    assert any(c.value == "Από το προφίλ σας" for c in cells)
    notes = [c for c in cells if isinstance(c.value, str) and "Στο προφίλ σας" in c.value]
    assert notes and all(c.data_type == "s" for c in notes)


def test_excel_keeps_its_shape_without_certificates(client, act, ai_on):
    _store()
    _customer(client)
    r = client.get(f"/act/{ADAM}/checklist.xlsx")
    ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
    assert not any(c.value == "Από το προφίλ σας"
                   for row in ws.iter_rows() for c in row)


# --------------------------------------------------------------------------- #
# Storing
# --------------------------------------------------------------------------- #
def test_saving_again_is_a_renewal_not_a_second_row(db):
    uid = make_user("certsave", "goodpassword1")
    cur = db.cursor()
    ev.save(cur, uid, scheme="iso9001", valid_until="2030-01-01")
    ev.save(cur, uid, scheme="iso9001", valid_until="2033-01-01", edition="2015")
    rows = ev.certificates(cur, uid)
    assert len(rows) == 1 and rows[0]["valid_until"] == dt.date(2033, 1, 1)
    # A manufacturer's is a separate row, one per manufacturer.
    ev.save(cur, uid, scheme="iso9001", holder="manufacturer", manufacturer="Α")
    ev.save(cur, uid, scheme="iso9001", holder="manufacturer", manufacturer="Β")
    assert len(ev.certificates(cur, uid)) == 3


@pytest.mark.parametrize("kw", [
    {"scheme": "iso10993"},
    {"scheme": "iso9001", "holder": "manufacturer"},       # needs a name
    {"scheme": "iso9001", "edition": "15"},
    {"scheme": "iso9001", "valid_until": "31/12/2030"},
    {"scheme": "iso9001", "number": "x" * 101},
])
def test_what_cannot_be_stored(db, kw):
    uid = make_user("certbad", "goodpassword1")
    with pytest.raises(ev.CertError):
        ev.save(db.cursor(), uid, **kw)


def _admin(client):
    make_user("certadmin", "goodpassword1", role="admin")
    login(client, "certadmin", "goodpassword1")
    return get_csrf(client)


def test_an_admin_declares_and_deletes_on_the_crm_card(client, db):
    uid = make_user("certcust", "goodpassword1")
    token = _admin(client)
    r = client.post(f"/admin/crm/{uid}/certificates",
                    data={"scheme": "iso14001", "holder": "self",
                          "valid_until": "2031-03-12", "edition": "2015",
                          "issuer": "TÜV", "csrf_token": token},
                    follow_redirects=False)
    assert r.status_code == 303 and "tab=fit" in r.headers["location"]
    card = client.get(f"/admin/crm/{uid}").text
    assert 'id="certs"' in card and "ISO 14001" in card and "12/03/2031" in card
    cid = ev.certificates(db.cursor(), uid)[0]["id"]
    client.post(f"/admin/crm/{uid}/certificates/{cid}/delete",
                data={"csrf_token": token})
    assert ev.certificates(db.cursor(), uid) == []


def test_a_customer_cannot_declare_through_the_crm(client, db):
    uid = _customer(client, "certself")
    r = client.post(f"/admin/crm/{uid}/certificates",
                    data={"scheme": "iso9001", "csrf_token": get_csrf(client)},
                    follow_redirects=False)
    assert r.status_code in (302, 303, 403, 404)
    assert ev.certificates(db.cursor(), uid) == []


# --------------------------------------------------------------------------- #
# Isolation (ai-summary §3)
# --------------------------------------------------------------------------- #
def _code(name):
    src = (ROOT / "app" / name).read_text()
    return re.sub(r'"""[\s\S]*?"""', "", src)


def test_the_summary_never_reads_certificates():
    code = _code("ai_summary.py")
    for forbidden in ("company_certificate", "eligibility_eval"):
        assert forbidden not in code


def test_the_evaluation_never_writes_the_summary():
    code = _code("eligibility_eval.py")
    assert "act_ai_summary" not in code
    assert "ai_summary" not in code
