"""
checklist_export.py — the checklist on paper and in a spreadsheet.

docs/specs/tender-checklist.md, slice 5: "printable / exportable list for the
person who assembles the envelope". That person is often not the one who
reads the notice on the site — a colleague, an accountant, the bank — so the
list has to leave the browser.

Two shapes of ONE thing, both built from tender_checklist.view(), the same
dict the panel renders, so paper, spreadsheet and screen cannot disagree:

  * /act/<adam>/checklist/print — a standalone page laid out for A4 (the
    browser's "Save as PDF" is the PDF export). Template:
    checklist_print.html. Nothing in this module; the route is in
    tender_checklist.make_router.
  * /act/<adam>/checklist.xlsx — workbook() below: the tasks on one sheet,
    the deadlines on another, real date cells so they sort.

Both carry the screening warning and the date they were made on: "5 days
left" printed on Monday is wrong on Wednesday, so a count is always stated
as "from <date>".

Safety: every text cell is written as a STRING. openpyxl turns any value
that starts with "=" into a formula, and these cells hold text quoted from a
notice and text the customer typed. workbook() never writes a formula.
"""
from __future__ import annotations

import datetime as dt
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from app import eligibility_eval as _eval
    from app import i18n as _i18n
except ImportError:                      # pragma: no cover — run with --app-dir=app
    import eligibility_eval as _eval
    import i18n as _i18n

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(vertical="top", wrap_text=True)


def disposition(adam: str, ext: str) -> str:
    """Content-Disposition for a download named after the act.

    A Diavgeia ΑΔΑ is Greek letters, which the plain filename= cannot carry:
    that part is an ASCII fallback, and filename* (RFC 5987) holds the real
    name for every browser that reads it."""
    from urllib.parse import quote
    ascii_name = "checklist-" + re.sub(r"[^A-Za-z0-9._-]", "_", adam) + "." + ext
    utf8_name = quote(f"checklist-{adam}.{ext}", safe="")
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{utf8_name}"


def _put(ws, row: int, col: int, value):
    """Write one cell; a string is ALWAYS stored as a string, never a formula."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _header(ws, row: int, labels: list[str], widths: list[int]) -> None:
    for i, (label, width) in enumerate(zip(labels, widths), start=1):
        cell = _put(ws, row, i, label)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _preamble(ws, cl: dict, meta: dict, t) -> int:
    """Title block at the top of a sheet. Returns the next free row."""
    made = meta["made_at"]
    lines = [
        (meta.get("title") or cl["adam"], Font(bold=True, size=13)),
        (f"ΑΔΑΜ {cl['adam']}" + (f" · {meta['authority']}" if meta.get("authority") else ""),
         Font(color="555555")),
        (t("Εξαγωγή") + " " + made.strftime("%d/%m/%Y %H:%M") + " · "
         + t("οι ημέρες μετρούν από αυτή την ημερομηνία."), Font(color="555555")),
        (t("Η λίστα προκύπτει αυτόματα από τη σύνοψη της προκήρυξης και μπορεί να είναι ελλιπής. Δεν αντικαθιστά την ανάγνωση της διακήρυξης· επιβεβαιώνετε πάντα τις προθεσμίες και τα δικαιολογητικά στα επίσημα έγγραφα."),
         Font(italic=True, color="8A4B00")),
    ]
    if cl.get("truncated"):
        lines.append((t("Η σύνοψη κάλυψε μόνο μέρος του κειμένου, οπότε η λίστα μπορεί να μην περιέχει όσα αναφέρονται στο υπόλοιπο."),
                      Font(italic=True, color="8A4B00")))
    for r, (text, font) in enumerate(lines, start=1):
        _put(ws, r, 1, text).font = font
    return len(lines) + 2


def _days(n) -> str:
    return "" if n is None else n


def workbook(cl: dict, meta: dict, lang: str = "el") -> bytes:
    """The checklist as .xlsx bytes. `cl` is tender_checklist.view()'s dict;
    `meta` = {"title", "authority", "made_at"}."""
    def t(s):
        return _i18n.translate(s, lang)

    wb = Workbook()

    # ---- sheet 1: what to prepare -------------------------------------- #
    ws = wb.active
    ws.title = t("Λίστα ελέγχου")[:31]
    row = _preamble(ws, cl, meta, t)
    # The evaluation layer's column exists only for a customer who declared
    # certificates — everyone else's workbook keeps its old shape.
    has_certs = any(item.get("certs") for g in cl["groups"] for item in g["items"])
    labels = [t("Ενότητα"), t("Στοιχείο"), t("Τι ζητείται"),
              t("Υποχρεωτικό"), t("Ολοκληρώθηκε"),
              t("Τι λέει η προκήρυξη"), t("Σημειώσεις")]
    widths = [24, 30, 40, 12, 14, 60, 30]
    if has_certs:
        labels.append(t("Από το προφίλ σας"))
        widths.append(44)
    ncol = len(labels)
    _header(ws, row, labels, widths)
    for g in cl["groups"]:
        for item in g["items"]:
            row += 1
            _put(ws, row, 1, t(g["heading"]))
            _put(ws, row, 2, item.get("label") or "")
            _put(ws, row, 3, item.get("value") or "")
            _put(ws, row, 4, t("υποχρεωτικό") if item.get("obligation") == "mandatory" else "")
            if item.get("done_at"):
                _put(ws, row, 5, item["done_at"].date()).number_format = "DD/MM/YYYY"
            _put(ws, row, 6, item.get("quote") or "")
            if item.get("certs"):
                _put(ws, row, 8, "\n".join(_eval.note_text(n, t) for n in item["certs"]))
            for col in range(1, ncol + 1):
                ws.cell(row=row, column=col).alignment = _WRAP
    for o in cl.get("own") or []:
        row += 1
        _put(ws, row, 1, t("Δικά σας"))
        _put(ws, row, 2, o["text"])
        if o.get("done_at"):
            _put(ws, row, 5, o["done_at"].date()).number_format = "DD/MM/YYYY"
        for col in range(1, ncol + 1):
            ws.cell(row=row, column=col).alignment = _WRAP

    # ---- sheet 2: by when ---------------------------------------------- #
    ws = wb.create_sheet(t("Προθεσμίες")[:31])
    row = _preamble(ws, cl, meta, t)
    _header(ws, row, [t("Ημερομηνία"), t("Ώρα"), t("Προθεσμία"),
                      t("Κείμενο"), t("Ημέρες"), t("Εργάσιμες ημέρες"), t("Πηγή")],
            [13, 8, 34, 40, 9, 11, 22])
    for d in cl["deadlines"]:
        row += 1
        if d["date"]:
            _put(ws, row, 1, d["date"]).number_format = "DD/MM/YYYY"
        _put(ws, row, 2, d.get("time") or "")
        label = t(d["label"]) if d["source"] == "record" else (d.get("label") or "")
        _put(ws, row, 3, label)
        _put(ws, row, 4, "" if d["date"] else (d.get("value") or ""))
        _put(ws, row, 5, _days(d.get("days_left")))
        _put(ws, row, 6, _days(d.get("workdays_left")))
        _put(ws, row, 7, t("από το αρχείο της πράξης") if d["source"] == "record"
             else t("σύνοψη AI"))
        for col in range(1, 8):
            ws.cell(row=row, column=col).alignment = _WRAP

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def lang_of(request) -> str:
    return _i18n.lang_from_request(request)


def made_at(tz) -> dt.datetime:
    return dt.datetime.now(tz).replace(microsecond=0)
